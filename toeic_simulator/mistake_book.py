"""
mistake_book.py
───────────────
Tracks questions the user has flagged as difficult / wrong.

Storage : mistake_book.json in BASE_DIR
Schema  : {"entries": [
              {"set_id": 1, "part": "Part1", "q_num": 1,
               "text": "...", "answer": "...", "ts": 1234567890},
              ...
           ]}
"""
import json
import os
import time
from typing import Optional


class MistakeBook:
    """Persistent store for marked questions."""

    def __init__(self, base_dir: str):
        self._path = os.path.join(base_dir, "mistake_book.json")
        self._entries: list = []
        self._load()

    # ── Public API ────────────────────────────────────────────────────────────

    def mark(self, set_id: int, part: str, q_num: int,
             text: str = "", answer: str = "") -> None:
        """Add a question; no-op if already present."""
        if self._find(set_id, part, q_num) is not None:
            return
        self._entries.append({
            "set_id": set_id,
            "part":   part,
            "q_num":  q_num,
            "text":   text,
            "answer": answer,
            "ts":     int(time.time()),
        })
        self._save()

    def unmark(self, set_id: int, part: str, q_num: int) -> None:
        """Remove a question from the mistake book."""
        idx = self._find(set_id, part, q_num)
        if idx is not None:
            self._entries.pop(idx)
            self._save()

    def toggle(self, set_id: int, part: str, q_num: int,
               text: str = "", answer: str = "") -> bool:
        """Toggle mark. Returns True if now marked, False if unmarked."""
        if self.is_marked(set_id, part, q_num):
            self.unmark(set_id, part, q_num)
            return False
        self.mark(set_id, part, q_num, text, answer)
        return True

    def is_marked(self, set_id: int, part: str, q_num: int) -> bool:
        return self._find(set_id, part, q_num) is not None

    def entries_for_set(self, set_id: int) -> list:
        return [e for e in self._entries if e["set_id"] == set_id]

    def all_entries(self) -> list:
        return list(self._entries)

    def count(self) -> int:
        return len(self._entries)

    def export_xlsx(self, out_path: str) -> str:
        """Export all mistakes to Excel. Returns saved path."""
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        from datetime import datetime

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "错题本"

        headers = ["套题编号", "题型", "题号", "题目 / 描述", "参考答案", "标记时间"]
        hdr_fill = PatternFill("solid", fgColor="003087")
        hdr_font = Font(bold=True, color="FFFFFF", name="Calibri")
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = hdr_font
            cell.fill = hdr_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        col_widths = [12, 10, 8, 60, 80, 20]
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

        for row_i, e in enumerate(self._entries, 2):
            ts_str = datetime.fromtimestamp(e.get("ts", 0)).strftime("%Y-%m-%d %H:%M")
            row_vals = [
                e.get("set_id", ""),
                e.get("part",   ""),
                e.get("q_num",  ""),
                e.get("text",   ""),
                e.get("answer", ""),
                ts_str,
            ]
            for col, val in enumerate(row_vals, 1):
                cell = ws.cell(row=row_i, column=col, value=val)
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        wb.save(out_path)
        return out_path

    # ── Private ───────────────────────────────────────────────────────────────

    def _find(self, set_id: int, part: str, q_num: int) -> Optional[int]:
        for i, e in enumerate(self._entries):
            if (e["set_id"] == set_id
                    and e["part"] == part
                    and e["q_num"] == q_num):
                return i
        return None

    def _load(self) -> None:
        if not os.path.isfile(self._path):
            return
        try:
            with open(self._path, encoding="utf-8") as f:
                data = json.load(f)
            self._entries = data.get("entries", [])
        except Exception as exc:
            print(f"[MistakeBook] Load error: {exc}")

    def _save(self) -> None:
        try:
            with open(self._path, "w", encoding="utf-8") as f:
                json.dump({"entries": self._entries}, f,
                          ensure_ascii=False, indent=2)
        except Exception as exc:
            print(f"[MistakeBook] Save error: {exc}")
